import io

import pytest
from docx import Document
from pypdf import PdfWriter

from texto_audio.extrair import ArquivoNaoSuportado, extrair_texto


def _docx(doc):
    destino = io.BytesIO()
    doc.save(destino)
    return destino.getvalue()


@pytest.mark.parametrize("codificacao", ["utf-8-sig", "utf-16", "utf-32", "cp1252"])
def test_texto_puro_respeita_codificacao(codificacao):
    texto = "Petição: ação e tutela."
    assert extrair_texto("peca.txt", texto.encode(codificacao)) == texto


def test_docx_preserva_ordem_das_tabelas_e_paragrafos():
    doc = Document()
    doc.add_paragraph("Antes da tabela.")
    tabela = doc.add_table(rows=1, cols=2)
    tabela.cell(0, 0).text = "Primeira célula."
    tabela.cell(0, 1).text = "Segunda célula."
    doc.add_paragraph("Depois da tabela.")
    assert extrair_texto("peca.docx", _docx(doc)) == (
        "Antes da tabela.\nPrimeira célula. Segunda célula.\nDepois da tabela.")


def test_docx_nao_repete_celula_mesclada():
    doc = Document()
    tabela = doc.add_table(rows=2, cols=2)
    tabela.cell(0, 0).merge(tabela.cell(1, 1)).text = "Uma célula mesclada."
    assert extrair_texto("peca.docx", _docx(doc)) == "Uma célula mesclada."


def test_docx_inclui_tabela_dentro_de_celula():
    doc = Document()
    celula = doc.add_table(rows=1, cols=1).cell(0, 0)
    celula.text = "Antes."
    celula.add_table(rows=1, cols=1).cell(0, 0).text = "Dentro."
    celula.add_paragraph("Depois.")
    assert extrair_texto("peca.docx", _docx(doc)) == "Antes. Dentro. Depois."


def test_pdf_protegido_tem_orientacao_clara():
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    pdf.encrypt("senha")
    dados = io.BytesIO()
    pdf.write(dados)
    with pytest.raises(ArquivoNaoSuportado, match="protegido por senha"):
        extrair_texto("peca.pdf", dados.getvalue())


def test_pdf_sem_texto_orienta_ocr():
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    dados = io.BytesIO()
    pdf.write(dados)
    with pytest.raises(ArquivoNaoSuportado, match="Faça OCR"):
        extrair_texto("peca.pdf", dados.getvalue())


def test_formato_antigo_orienta_docx():
    with pytest.raises(ArquivoNaoSuportado, match="Salve como .docx"):
        extrair_texto("peca.doc", b"")
