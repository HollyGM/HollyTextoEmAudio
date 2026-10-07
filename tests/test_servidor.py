import asyncio
import io
import shutil
import threading
import time

import pytest
from fastapi.testclient import TestClient

from texto_audio import pipeline, servidor

from .conftest import MotorFalso, MotorQuebrado

pytestmark = pytest.mark.skipif(not shutil.which("ffmpeg"),
                                reason="ffmpeg não instalado")


@pytest.fixture
def cliente(tmp_path, monkeypatch):
    monkeypatch.setattr(servidor, "PASTA_SAIDAS", tmp_path)
    monkeypatch.setattr(servidor, "ARQUIVO_PRONUNCIAS", tmp_path / "nao-existe.txt")
    monkeypatch.setattr(pipeline, "MotorEdge", MotorFalso)
    servidor.TAREFAS.clear()
    with TestClient(servidor.app, base_url="http://127.0.0.1") as c:
        yield c


def esperar(cliente, tarefa_id, limite=15):
    fim = time.time() + limite
    while time.time() < fim:
        estado = cliente.get(f"/api/audio/{tarefa_id}").json()
        if estado["status"] in ("pronto", "erro"):
            return estado
        time.sleep(0.1)
    raise AssertionError("tarefa não terminou")


def test_pagina_e_vozes(cliente):
    assert "Texto em Áudio" in cliente.get("/").text
    vozes = cliente.get("/api/vozes").json()
    assert vozes["padrao"] in [v["id"] for v in vozes["vozes"]]
    assert cliente.get("/api/vozes").headers["cache-control"] == "no-store"


def test_info_devolve_pasta_real_e_quantidade_de_tarefas(cliente, tmp_path):
    info = cliente.get("/api/info").json()
    assert info["status"] == "pronto"
    assert info["versao"] == servidor.__version__
    assert info["pasta_saidas"] == str(tmp_path)
    assert info["tarefas_ativas"] == 0


@pytest.mark.parametrize("cabecalhos", [
    {"Host": "site-externo.example"},
    {"Origin": "https://site-externo.example"},
    {"Origin": "null"},
    {"Origin": "http://127.0.0.1:8766"},
    {"Sec-Fetch-Site": "cross-site"},
])
def test_acesso_de_site_externo_e_rejeitado(cliente, cabecalhos):
    assert cliente.get("/api/info", headers=cabecalhos).status_code == 403
    assert cliente.post("/api/audio", json={"texto": "Olá"},
                        headers=cabecalhos).status_code == 403
    assert not servidor.TAREFAS


def test_origem_local_da_interface_e_aceita(cliente):
    assert cliente.post("/api/preparar", json={"texto": "Olá"},
                        headers={"Origin": "http://127.0.0.1"}).status_code == 200


def _conversao_bloqueada(monkeypatch, tmp_path):
    """Controla a liberação sem rede nem processamento de áudio."""
    liberar = threading.Event()
    iniciadas = []

    async def gerar(texto, pasta, cfg, progresso):
        iniciadas.append(texto)
        while not liberar.is_set():
            await asyncio.sleep(0.01)
        caminho = tmp_path / f"audio-{len(iniciadas)}.mp3"
        caminho.write_bytes(b"mp3")
        return pipeline.Resultado(caminho, 1.0, 3, 1, 1)

    monkeypatch.setattr(servidor, "gerar_audio", gerar)
    return liberar, iniciadas


def test_conversoes_aguardam_na_fila_sem_serem_descartadas(cliente, tmp_path, monkeypatch):
    liberar, iniciadas = _conversao_bloqueada(monkeypatch, tmp_path)
    primeiro = cliente.post("/api/audio", json={"texto": "Primeiro"}).json()["id"]
    segundo = cliente.post("/api/audio", json={"texto": "Segundo"}).json()["id"]
    assert cliente.get(f"/api/audio/{primeiro}").json()["status"] == "gerando"
    fila = cliente.get(f"/api/audio/{segundo}").json()
    assert fila["status"] == "na_fila" and fila["posicao_fila"] == 2
    assert iniciadas == ["Primeiro"]

    # A tarefa ativa é a mais antiga: a limpeza deve pular até uma concluída.
    monkeypatch.setattr(servidor, "MAX_TAREFAS_GUARDADAS", 3)
    servidor.TAREFAS["concluida"] = servidor.Tarefa(status="pronto")
    terceiro = cliente.post("/api/audio", json={"texto": "Terceiro"}).json()["id"]
    assert "concluida" not in servidor.TAREFAS
    assert cliente.get(f"/api/audio/{primeiro}").status_code == 200
    assert cliente.get("/api/info").json()["tarefas_ativas"] == 3
    liberar.set()
    for tarefa_id in (primeiro, segundo, terceiro):
        assert esperar(cliente, tarefa_id)["status"] == "pronto"
    assert iniciadas == ["Primeiro", "Segundo", "Terceiro"]


def test_fila_cheia_recusa_sem_perder_tarefa(cliente, tmp_path, monkeypatch):
    liberar, _ = _conversao_bloqueada(monkeypatch, tmp_path)
    monkeypatch.setattr(servidor, "MAX_TAREFAS_PENDENTES", 1)
    primeiro = cliente.post("/api/audio", json={"texto": "Primeiro"}).json()["id"]
    resposta = cliente.post("/api/audio", json={"texto": "Segundo"})
    assert resposta.status_code == 429
    assert list(servidor.TAREFAS) == [primeiro]
    liberar.set()
    assert esperar(cliente, primeiro)["status"] == "pronto"


def test_falha_de_configuracao_nao_deixa_geracao_travada(cliente, monkeypatch):
    def quebrar(_):
        raise OSError("pronúncias inacessíveis")

    monkeypatch.setattr(servidor, "_opcoes", quebrar)
    tarefa_id = cliente.post("/api/audio", json={"texto": "Olá"}).json()["id"]
    estado = esperar(cliente, tarefa_id)
    assert estado["status"] == "erro"
    assert "Reinicie" in estado["erro"]


def test_falha_de_montagem_tem_mensagem_legivel(cliente, monkeypatch):
    async def quebrar(*args, **kwargs):
        raise servidor.ErroAudio("ffmpeg não encontrado")

    monkeypatch.setattr(servidor, "gerar_audio", quebrar)
    tarefa_id = cliente.post("/api/audio", json={"texto": "Olá"}).json()["id"]
    estado = esperar(cliente, tarefa_id)
    assert estado["status"] == "erro"
    assert estado["erro"] == "Não consegui montar o MP3: ffmpeg não encontrado"


def test_fechamento_cancela_conversao_pendente(tmp_path, monkeypatch):
    _conversao_bloqueada(monkeypatch, tmp_path)
    servidor.TAREFAS.clear()
    with TestClient(servidor.app, base_url="http://127.0.0.1") as cliente:
        tarefa_id = cliente.post("/api/audio", json={"texto": "Olá"}).json()["id"]
        tarefa = servidor.TAREFAS[tarefa_id]
    assert tarefa.status == "erro"
    assert "interrompida" in tarefa.erro
    assert tarefa._tarefa.cancelled()


def test_preparar_devolve_texto_falado_e_estimativa(cliente):
    r = cliente.post("/api/preparar", json={
        "texto": "Conforme o art. 5º da CF/88.", "velocidade": 0}).json()
    assert "artigo quinto" in r["texto_falado"]
    assert r["palavras"] > 0 and r["minutos"] > 0 and r["blocos"] == 1
    mais_rapido = cliente.post("/api/preparar", json={
        "texto": "Conforme o art. 5º da CF/88.", "velocidade": 50}).json()
    assert mais_rapido["minutos"] < r["minutos"]


def test_pronuncias_da_requisicao_valem(cliente):
    r = cliente.post("/api/preparar", json={
        "texto": "Pelo BNDES.", "pronuncias": "BNDES = bê ene dê é ésse"}).json()
    assert r["texto_falado"] == "Pelo bê ene dê é ésse."


def test_fluxo_completo_gera_e_serve_o_mp3(cliente, tmp_path):
    resp = cliente.post("/api/audio", json={"texto": "Olá.\nMundo.", "titulo": "Teste"})
    assert resp.status_code == 202
    estado = esperar(cliente, resp.json()["id"])
    assert estado["status"] == "pronto", estado
    assert estado["duracao_s"] > 0 and estado["arquivo"].endswith(".mp3")

    audio = cliente.get(f"/api/audio/{resp.json()['id']}/arquivo")
    assert audio.status_code == 200
    assert audio.headers["content-type"] == "audio/mpeg"
    assert len(audio.content) == estado["tamanho_bytes"]
    assert (tmp_path / estado["arquivo"]).exists()


def test_falha_da_voz_vira_erro_legivel(cliente, monkeypatch):
    monkeypatch.setattr(pipeline, "MotorEdge", MotorQuebrado)
    tarefa_id = cliente.post("/api/audio", json={"texto": "Olá."}).json()["id"]
    estado = esperar(cliente, tarefa_id)
    assert estado["status"] == "erro" and "fora do ar" in estado["erro"]
    assert cliente.get(f"/api/audio/{tarefa_id}/arquivo").status_code == 404


@pytest.mark.parametrize("corpo", [
    {"texto": "   "},
    {"texto": "Olá", "voz": "voz-que-nao-existe"},
    {"texto": "Olá", "velocidade": 500},
    {"texto": "Olá", "siglas": "gritando"},
    {"texto": "x" * (servidor.LIMITE_TEXTO + 1)},
])
def test_pedidos_invalidos_sao_recusados(cliente, corpo):
    assert cliente.post("/api/audio", json=corpo).status_code == 422


def test_tarefa_inexistente(cliente):
    assert cliente.get("/api/audio/naoexiste").status_code == 404
    assert cliente.get("/api/audio/naoexiste/arquivo").status_code == 404


def test_mp3_excluido_tem_mensagem_legivel(cliente, tmp_path):
    servidor.TAREFAS["apagada"] = servidor.Tarefa(
        status="pronto", caminho=tmp_path / "apagado.mp3")
    resposta = cliente.get("/api/audio/apagada/arquivo")
    assert resposta.status_code == 404
    assert "movido ou excluído" in resposta.json()["detail"]


def test_extrair_txt_docx_e_formatos_invalidos(cliente):
    txt = cliente.post("/api/extrair", files={
        "arquivo": ("a.txt", "Ação de cobrança".encode("cp1252"))})
    assert txt.json()["texto"] == "Ação de cobrança"

    from docx import Document
    doc = Document()
    doc.add_paragraph("Primeiro parágrafo.")
    doc.add_paragraph("Segundo parágrafo.")
    buf = io.BytesIO(); doc.save(buf)
    r = cliente.post("/api/extrair", files={"arquivo": ("p.docx", buf.getvalue())})
    assert r.json()["texto"] == "Primeiro parágrafo.\nSegundo parágrafo."

    assert cliente.post("/api/extrair", files={
        "arquivo": ("p.doc", b"x")}).status_code == 415
    assert cliente.post("/api/extrair", files={
        "arquivo": ("p.exe", b"x")}).status_code == 415
    assert cliente.post("/api/extrair", files={
        "arquivo": ("quebrado.docx", b"isto nao e um docx")}).status_code == 422


def test_upload_vazio_e_grande_sao_recusados(cliente, monkeypatch):
    assert cliente.post("/api/extrair", files={
        "arquivo": ("vazio.txt", b" \n ")}).status_code == 422
    monkeypatch.setattr(servidor, "LIMITE_UPLOAD", 10)
    assert cliente.post("/api/extrair", files={
        "arquivo": ("grande.txt", b"x" * 11)}).status_code == 413


def test_upload_longo_informa_truncamento(cliente, monkeypatch):
    monkeypatch.setattr(servidor, "LIMITE_TEXTO", 5)
    resposta = cliente.post("/api/extrair", files={
        "arquivo": ("longo.txt", b"123456789")})
    assert resposta.status_code == 200
    assert resposta.json() == {"texto": "12345", "truncado": True}


def _pdf_minimo(texto: str) -> bytes:
    """PDF de uma página com `texto` (xref correto) para testar a extração."""
    conteudo = f"BT /F1 14 Tf 72 720 Td ({texto}) Tj ET".encode("latin-1")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(conteudo) + conteudo + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
        b"/Encoding /WinAnsiEncoding >>",
    ]
    saida, offsets = b"%PDF-1.4\n", []
    for i, obj in enumerate(objetos, 1):
        offsets.append(len(saida))
        saida += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(saida)
    saida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
    for off in offsets:
        saida += b"%010d 00000 n \n" % off
    saida += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
              % (len(objetos) + 1, xref))
    return saida


def test_extrair_pdf_com_texto(cliente):
    r = cliente.post("/api/extrair", files={
        "arquivo": ("peca.pdf", _pdf_minimo("Acao de cobranca"))})
    assert r.status_code == 200
    assert "Acao de cobranca" in r.json()["texto"]


def test_pdf_sem_texto_explica_que_e_imagem(cliente):
    r = cliente.post("/api/extrair", files={
        "arquivo": ("scan.pdf", _pdf_minimo(""))})
    assert r.status_code == 415 and "OCR" in r.json()["detail"]
